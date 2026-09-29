"""The resume skill teaches code the agent writes and runs; that code must work.

Runs the skill's helper patterns + example script exactly as written in
SKILL.md, then checks the three outputs with independent logic (`docx_samples`):

- review: tracked changes; Reject All restores the original text AND formatting
- redline: the same changes as plain colour formatting, no tracking
- final: every change applied, clean (no markup, comments or review colours)

plus the colour of each kind of change, LibreOffice opening every file, and the
same script running inside the Docker sandbox.
"""

from __future__ import annotations

import shutil
import subprocess
import zipfile

import pytest
from genai_agentic_sandbox.builder import SandboxImageBuilder
from genai_agentic_sandbox.main import SKILLS_DIR
from genai_agentic_sandbox.sandbox.docker import Mount

from .docx_samples import (
    PALETTE,
    RESUME,
    W,
    comments,
    part,
    render,
    revision_ids,
    revisions,
    runs_info,
    skill_code,
)

SKILL = (SKILLS_DIR / "resume" / "docx-tracked-revisions" / "SKILL.md").read_text()
HELPERS = skill_code(SKILL, "tracked-change helpers")
EXAMPLE = skill_code(SKILL, "example")
ORIGINAL = render(RESUME, "original")

EXPECTED = [
    "Jane Doe",
    "Backend engineer: 5 years building Python REST APIs.",  # added summary
    "Senior Python engineer with 5 years of experience building APIs.",  # rephrased
    "Skills",
    "Docker",  # moved up
    "Python, Django, FastAPI, SQL",  # added keyword
    "Experience",
    "Acme Corp — Backend Engineer (2020–2025)",  # styled + flagged, text unchanged
    "Built REST services.",  # clause removed
    "Led a team of 3 engineers; cut latency 40%.",  # added bullet
    # "References available on request." removed
    "Email",
    "jane@example.com",
]


def script_for(src: str, job_dir: str) -> str:
    example = EXAMPLE.replace('"/input/cv.docx"', repr(src)).replace(
        '"/output/resume/1-acme-backend"', repr(job_dir)
    )
    return f"{HELPERS}\n\n{example}\n"


@pytest.fixture(scope="module")
def outputs(tmp_path_factory) -> dict:
    """Run the skill's code on the host Python (python-docx is a dev dependency)."""
    d = tmp_path_factory.mktemp("skill")
    (d / "cv.docx").write_bytes(RESUME)
    script = d / "revise_resume.py"
    script.write_text(script_for(str(d / "cv.docx"), str(d / "job")))
    run = subprocess.run(
        ["python", str(script)], capture_output=True, text=True, check=False, timeout=120
    )
    assert run.returncode == 0, run.stderr
    v1 = d / "job" / "v1"
    return {
        "stdout": run.stdout,
        "review": (v1 / "cv_acme-backend_review.docx").read_bytes(),
        "redline": (v1 / "cv_acme-backend_redline.docx").read_bytes(),
        "final": (v1 / "cv_acme-backend_final.docx").read_bytes(),
    }


def by_text(info: list[dict], text: str) -> list[dict]:
    return [r for r in info if r["text"] == text]


# -- the script's own report -------------------------------------------------------


def test_every_change_ok_and_every_check_true(outputs):
    lines = outputs["stdout"].splitlines()
    status = [line for line in lines if line.startswith(("OK", "FAIL"))]
    assert len(status) == 11 and all(s.startswith("OK") for s in status), status
    checks = [line for line in lines if line.startswith("CHECK")]
    assert len(checks) == 4 and all(c.endswith("True") for c in checks), checks


# -- review copy (tracked) -----------------------------------------------------------


def test_review_reject_all_restores_the_original_text(outputs):
    assert render(outputs["review"], "original") == ORIGINAL


def test_review_accept_all_gives_the_tailored_resume(outputs):
    assert render(outputs["review"], "accepted") == EXPECTED


def test_review_markup_covers_every_kind_of_change(outputs):
    counts = revisions(outputs["review"])
    assert counts["ins"] >= 4 and counts["del"] >= 3
    assert counts["moveFrom"] == counts["moveTo"] == 1
    assert counts["rPrChange"] >= 5  # format change + colour on existing text
    assert counts["marks"] == 5  # 2 new paragraphs, 1 deleted, move from/to
    assert len(part(outputs["review"]).findall(f".//{W}pPrChange")) == 1  # style change


def test_review_colour_on_existing_text_is_reversible(outputs):
    """Reject All must restore formatting too: recoloured original text keeps its old rPr."""
    for run in runs_info(outputs["review"]):
        if run["color"] in PALETTE.values() and run["wrapper"] not in ("ins", "moveTo"):
            assert run["has_change"], run
            assert run["old_color"] not in PALETTE.values(), run


def test_review_comments_start_with_the_legend(outputs):
    notes = comments(outputs["review"])
    assert notes[0].startswith("Review colours - green: added")
    assert "Posting title: Senior Python Engineer" in notes
    assert any(
        n.startswith("Possible error: role end date is in the future") and "2020–Present" in n
        for n in notes
    )
    assert len(notes) == 6


def test_review_ids_unique_and_track_changes_on(outputs):
    ids = revision_ids(outputs["review"])
    assert len(ids) == len(set(ids))
    assert part(outputs["review"], "word/settings.xml").find(f"{W}trackRevisions") is not None


# -- colours -------------------------------------------------------------------------


@pytest.mark.parametrize("copy", ["review", "redline"])
def test_colour_per_kind_of_change(outputs, copy):
    info = runs_info(outputs[copy])
    assert (
        by_text(info, "Backend engineer: 5 years building Python REST APIs.")[0]["color"]
        == PALETTE["added"]
    )
    assert by_text(info, ", FastAPI")[0]["color"] == PALETTE["added"]
    assert by_text(info, "Python developer")[0]["color"] == PALETTE["removed"]
    assert by_text(info, "Senior Python engineer")[0]["color"] == PALETTE["rephrased"]
    assert by_text(info, "References available on request.")[0]["color"] == PALETTE["removed"]
    moved = by_text(info, "Docker")
    assert len(moved) == 2 and all(r["color"] == PALETTE["moved"] for r in moved)
    flag = by_text(info, "2020–2025")[0]
    assert (
        flag["color"] == PALETTE["flag"]
        and flag["fill"] == "FFC7CE"
        and flag["underline"] == "double"
    )
    assert by_text(info, "Acme Corp")[0]["fill"] == "FFF2CC"  # styling change


def test_redline_has_no_tracking_and_strikes_removed_text(outputs):
    content = outputs["redline"]
    counts = revisions(content)
    assert counts["ins"] == counts["del"] == counts["moveFrom"] == counts["moveTo"] == 0
    assert counts["rPrChange"] == 0 and counts["marks"] == 0
    info = runs_info(content)
    assert by_text(info, "Python developer")[0]["strike"] == "true"
    assert by_text(info, "References available on request.")[0]["strike"] == "true"
    assert len([r for r in by_text(info, "Docker") if r["strike"]]) == 1  # moved-away copy
    assert by_text(info, "Senior Python engineer")[0]["strike"] is None
    assert "References available on request." in render(content)  # removed text still readable
    assert comments(content)[0].startswith("Review colours")
    assert part(content, "word/settings.xml").find(f"{W}trackRevisions") is None


# -- final copy ----------------------------------------------------------------------


def test_final_is_the_tailored_resume_and_clean(outputs):
    content = outputs["final"]
    assert render(content, "original") == render(content, "accepted") == EXPECTED
    body = part(content)
    for tag in (
        "ins",
        "del",
        "moveFrom",
        "moveTo",
        "rPrChange",
        "pPrChange",
        "commentReference",
        "commentRangeStart",
        "moveFromRangeStart",
    ):
        assert body.find(f".//{W}{tag}") is None, tag
    for run in runs_info(content):
        assert run["color"] not in PALETTE.values(), run
        assert run["fill"] is None and run["strike"] is None, run
    assert part(content, "word/settings.xml").find(f"{W}trackRevisions") is None


def test_final_keeps_the_intended_formatting(outputs):
    paragraphs = list(part(outputs["final"]).iter(f"{W}p"))
    acme = next(p for p in paragraphs if "Acme Corp" in "".join(p.itertext()))
    assert acme.find(f"{W}pPr/{W}pStyle").get(f"{W}val") == "Heading2"  # style change applied
    bold = [r for r in acme.iter(f"{W}r") if r.find(f"{W}rPr/{W}b") is not None]
    assert "".join("".join(r.itertext()) for r in bold) == "Acme Corp"
    assert paragraphs[1].find(f"{W}pPr/{W}pStyle") is None  # summary in body style (like=P[1])


# -- other patterns ------------------------------------------------------------------


def test_find_failure_is_reported_not_fatal(tmp_path):
    (tmp_path / "cv.docx").write_bytes(RESUME)
    script = tmp_path / "s.py"
    script.write_text(
        HELPERS
        + f"""
doc = open_doc({str(tmp_path / "cv.docx")!r})
P = paragraphs(doc)
try:
    track_delete(P[1], "COBOL")
except ValueError as exc:
    print("FAIL", exc)
track_insert(P[4], ", Kubernetes")
doc.save({str(tmp_path / "out.docx")!r})
"""
    )
    run = subprocess.run(["python", str(script)], capture_output=True, text=True, check=False)
    assert run.returncode == 0, run.stderr
    assert "'COBOL' not found in" in run.stdout
    assert render((tmp_path / "out.docx").read_bytes())[4] == "Docker, Kubernetes"


def test_two_tailored_copies_from_one_input(tmp_path):
    (tmp_path / "cv.docx").write_bytes(RESUME)
    for rank, slug, title in (
        (1, "acme-backend", "Backend Engineer"),
        (2, "beta-platform", "Platform Engineer"),
    ):
        out = tmp_path / f"{rank}-{slug}" / f"cv_{slug}.docx"
        script = tmp_path / f"job{rank}.py"
        script.write_text(
            HELPERS
            + f"""
from pathlib import Path
Path({str(out)!r}).parent.mkdir(parents=True, exist_ok=True)
doc = open_doc({str(tmp_path / "cv.docx")!r})
P = paragraphs(doc)
d, i = track_replace(P[1], "Python developer", {title!r})
add_comment(doc, P[1], d, i, "Title of job {rank}")
enable_track_changes(doc)
doc.save({str(out)!r})
"""
        )
        run = subprocess.run(["python", str(script)], capture_output=True, text=True, check=False)
        assert run.returncode == 0, run.stderr
    one = (tmp_path / "1-acme-backend" / "cv_acme-backend.docx").read_bytes()
    two = (tmp_path / "2-beta-platform" / "cv_beta-platform.docx").read_bytes()
    assert render(one)[1].startswith("Backend Engineer with 5 years")
    assert render(two)[1].startswith("Platform Engineer with 5 years")
    assert render(one, "original") == render(two, "original") == ORIGINAL
    assert (tmp_path / "cv.docx").read_bytes() == RESUME


def test_deleting_a_field_code_marks_instr_text_deleted(tmp_path):
    import docx
    from docx.oxml import OxmlElement

    d = docx.Document()
    p = d.add_paragraph("Portfolio: ")
    for kind, text in (("begin", None), (None, ' HYPERLINK "https://x.dev" '), ("separate", None)):
        r = OxmlElement("w:r")
        if kind:
            r.append(OxmlElement("w:fldChar", attrs={f"{W}fldCharType": kind}))
        else:
            it = OxmlElement("w:instrText")
            it.text = text
            r.append(it)
        p._p.append(r)
    p.add_run("x.dev")
    end = OxmlElement("w:r")
    end.append(OxmlElement("w:fldChar", attrs={f"{W}fldCharType": "end"}))
    p._p.append(end)
    src = tmp_path / "f.docx"
    d.save(src)
    script = tmp_path / "s.py"
    script.write_text(
        HELPERS
        + f"""
doc = open_doc({str(src)!r})
P = paragraphs(doc)
delete_paragraph(P[0])
doc.save({str(tmp_path / "out.docx")!r})
accept_all(doc)
doc.save({str(tmp_path / "final.docx")!r})
"""
    )
    run = subprocess.run(["python", str(script)], capture_output=True, text=True, check=False)
    assert run.returncode == 0, run.stderr
    body = part((tmp_path / "out.docx").read_bytes())
    assert body.find(f".//{W}instrText") is None
    assert body.find(f".//{W}del//{W}delInstrText") is not None
    assert render((tmp_path / "out.docx").read_bytes(), "accepted") == []
    assert render((tmp_path / "final.docx").read_bytes()) == []


def test_move_then_edit_the_moved_copy(tmp_path):
    (tmp_path / "cv.docx").write_bytes(RESUME)
    script = tmp_path / "s.py"
    script.write_text(
        HELPERS
        + f"""
doc = open_doc({str(tmp_path / "cv.docx")!r})
P = paragraphs(doc)
moved = move_paragraph_after(P[4], P[2], "move1")
track_insert(moved, ", Kubernetes")
d, i = track_replace(moved, "Docker", "Docker Compose")
add_comment(doc, moved, d, i, "Named in the posting")
doc.save({str(tmp_path / "out.docx")!r})
"""
    )
    run = subprocess.run(["python", str(script)], capture_output=True, text=True, check=False)
    assert run.returncode == 0, run.stderr
    content = (tmp_path / "out.docx").read_bytes()
    counts = revisions(content)
    assert counts["moveFrom"] == counts["moveTo"] == 1
    assert render(content, "original") == ORIGINAL
    accepted = render(content, "accepted")
    assert accepted[3] == "Docker Compose, Kubernetes"
    assert accepted[4] == "Python, Django, SQL"


# -- real word processor & sandbox ------------------------------------------------------


@pytest.mark.skipif(shutil.which("soffice") is None, reason="LibreOffice not installed")
@pytest.mark.parametrize("copy", ["review", "redline", "final"])
def test_libreoffice_opens_every_copy(outputs, copy, tmp_path):
    src = tmp_path / f"{copy}.docx"
    src.write_bytes(outputs[copy])
    run = subprocess.run(
        [
            "soffice",
            f"-env:UserInstallation=file://{tmp_path}/lo",
            "--headless",
            "--convert-to",
            "odt",
            "--outdir",
            str(tmp_path),
            str(src),
        ],
        capture_output=True,
        text=True,
        timeout=180,
        check=False,
    )
    assert run.returncode == 0, run.stderr
    with zipfile.ZipFile(tmp_path / f"{copy}.odt") as zf:
        xml = zf.read("content.xml").decode()
    assert ("tracked-changes" in xml) == (copy == "review")
    assert ("JobHunter AI" in xml) == (copy != "final")  # comment / revision author


@pytest.mark.integration
def test_same_script_runs_inside_the_sandbox(tmp_path):
    """What the agent actually does: write the script into /output and `execute` it."""
    resume = tmp_path / "cv.docx"
    resume.write_bytes(RESUME)
    out = tmp_path / "out"
    out.mkdir()
    backend = SandboxImageBuilder().backend(
        mounts=[
            Mount(resume, "/input/cv.docx", read_only=True),
            Mount(out, "/output", read_only=False),
        ]
    )
    job = "/output/resume/1-acme-backend"
    with backend:
        script = script_for("/input/cv.docx", job)
        assert backend.write(f"{job}/revise_resume.py", script).error is None
        result = backend.execute(f"python {job}/revise_resume.py")
    assert result.exit_code == 0, result.output
    assert result.output.count("OK ") == 11
    assert "False" not in result.output
    v1 = out / "resume" / "1-acme-backend" / "v1"
    assert render((v1 / "cv_acme-backend_review.docx").read_bytes(), "original") == ORIGINAL
    assert render((v1 / "cv_acme-backend_final.docx").read_bytes()) == EXPECTED
    assert (v1 / "cv_acme-backend_redline.docx").exists()
    assert resume.read_bytes() == RESUME
