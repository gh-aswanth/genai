"""Test-only helpers: a sample resume and an independent reader of tracked changes.

Deliberately separate from the skill's helper code, so tests check what the
agent's scripts produce with logic the scripts do not share.
"""

from __future__ import annotations

import io
import re
import zipfile

import docx
from lxml import etree

W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def make_resume() -> bytes:
    """Multi-run paragraphs, a bold run, bullets, a tab and a table."""
    d = docx.Document()
    d.add_heading("Jane Doe", level=1)  # 0
    p = d.add_paragraph()  # 1
    p.add_run("Python developer with ")
    p.add_run("5 years").bold = True
    p.add_run(" of experience building APIs.")
    d.add_heading("Skills", level=2)  # 2
    d.add_paragraph("Python, Django, SQL", style="List Bullet")  # 3
    d.add_paragraph("Docker", style="List Bullet")  # 4
    d.add_heading("Experience", level=2)  # 5
    d.add_paragraph("Acme Corp — Backend Engineer (2020–2025)")  # 6
    d.add_paragraph("Built REST services.\tLed a team of 3.")  # 7
    d.add_paragraph("References available on request.")  # 8
    table = d.add_table(rows=1, cols=2)
    table.cell(0, 0).text = "Email"  # 9
    table.cell(0, 1).text = "jane@example.com"  # 10
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


RESUME = make_resume()


def part(content: bytes, name: str = "word/document.xml"):
    with zipfile.ZipFile(io.BytesIO(content)) as zf:
        return etree.fromstring(zf.read(name))


def render(content: bytes, view: str = "accepted") -> list[str]:
    """Paragraph texts with every change accepted, or every change rejected."""
    hide = {f"{W}del", f"{W}moveFrom"} if view == "accepted" else {f"{W}ins", f"{W}moveTo"}
    lines = []
    for p in part(content).iter(f"{W}p"):
        mark = p.find(f"{W}pPr/{W}rPr")
        gone = mark is not None and any(c.tag in hide for c in mark)
        text = ""
        for el in p.iter(f"{W}t", f"{W}delText", f"{W}tab"):
            if any(a.tag in hide for a in el.iterancestors()):
                continue
            text += "\t" if el.tag == f"{W}tab" else (el.text or "")
        if not (gone and not text):
            lines.append(text)
    return lines


def revisions(content: bytes) -> dict[str, int]:
    """Content-level revision counts; paragraph-mark flags counted separately."""
    body = part(content)
    ns = {"w": W[1:-1]}
    counts = {
        tag: len(body.xpath(f".//w:{tag}[not(parent::w:rPr/parent::w:pPr)]", namespaces=ns))
        for tag in ("ins", "del", "moveFrom", "moveTo", "rPrChange")
    }
    counts["marks"] = len(
        body.xpath(
            ".//w:pPr/w:rPr/*[self::w:ins or self::w:del or self::w:moveFrom or self::w:moveTo]",
            namespaces=ns,
        )
    )
    counts["comments"] = len(body.xpath(".//w:commentReference", namespaces=ns))
    return counts


def comments(content: bytes) -> list[str]:
    return ["".join(c.itertext()) for c in part(content, "word/comments.xml").iter(f"{W}comment")]


def revision_ids(content: bytes) -> list[str]:
    tags = (
        "ins",
        "del",
        "moveFrom",
        "moveTo",
        "rPrChange",
        "moveFromRangeStart",
        "moveToRangeStart",
    )
    body = part(content)
    return [e.get(f"{W}id") for t in tags for e in body.iter(f"{W}{t}")]


def skill_code(skill_md: str, marker: str) -> str:
    """The ```python block of a SKILL.md whose first line is `# --- <marker> ---`."""
    for block in re.findall(r"^```python\n(.*?)\n```$", skill_md, re.DOTALL | re.MULTILINE):
        if block.startswith(f"# --- {marker} ---"):
            return block
    raise LookupError(f"no python block marked {marker!r}")


PALETTE = {
    "added": "00B050",
    "removed": "FF0000",
    "rephrased": "0070C0",
    "moved": "7030A0",
    "flag": "C00000",
}


def runs_info(content: bytes) -> list[dict]:
    """Every text run: its text, colour, shading, strike, underline and revision wrapper."""
    out = []
    for r in part(content).iter(f"{W}r"):
        text = "".join(t.text or "" for t in r.iter(f"{W}t", f"{W}delText"))
        if not text:
            continue
        rpr = r.find(f"{W}rPr")

        def val(tag, attr="val", rpr=rpr):
            el = rpr.find(f"{W}{tag}") if rpr is not None else None
            return None if el is None else el.get(f"{W}{attr}", "true")

        wrapper = next(
            (
                a.tag[len(W) :]
                for a in r.iterancestors()
                if a.tag in {f"{W}{t}" for t in ("ins", "del", "moveFrom", "moveTo")}
            ),
            None,
        )
        change = rpr.find(f"{W}rPrChange") if rpr is not None else None
        old = change.find(f"{W}rPr") if change is not None else None
        out.append(
            {
                "text": text,
                "color": val("color"),
                "fill": val("shd", "fill"),
                "strike": val("strike"),
                "underline": val("u"),
                "wrapper": wrapper,
                "old_color": None
                if old is None or old.find(f"{W}color") is None
                else old.find(f"{W}color").get(f"{W}val"),
                "has_change": change is not None,
            }
        )
    return out
