"""The orchestrator's skills: routing, cleanup and the ATS-only report must be sound."""

from __future__ import annotations

import io
import json
import os
import re
import subprocess

import docx
import pytest
from genai_agentic_sandbox.main import SKILLS_DIR

from .docx_samples import skill_code

ORCH = SKILLS_DIR / "orchestrator"
EXPECTED_SKILLS = {
    "agent-orchestration",
    "jobhunt-workflow",
    "resume-optimization",
    "ats-score-only",
    "job-posting-intake",
    "output-cleanup",
    "agent-memory",
    "workflow-planning",
}
CLEANUP = skill_code((ORCH / "output-cleanup" / "SKILL.md").read_text(), "cleanup reference")
REPORT = skill_code((ORCH / "ats-score-only" / "SKILL.md").read_text(), "ats-report reference")


def test_orchestrator_has_all_workflow_skills():
    assert {p.parent.name for p in ORCH.glob("*/SKILL.md")} == EXPECTED_SKILLS


def test_routing_table_covers_every_workflow():
    routing = (ORCH / "agent-orchestration" / "SKILL.md").read_text()
    table = routing[routing.index("## 1. Route") : routing.index("## 2.")]
    for skill in EXPECTED_SKILLS - {"agent-orchestration"}:
        assert f"`{skill}`" in table, skill


def test_score_only_never_edits():
    text = (ORCH / "ats-score-only" / "SKILL.md").read_text()
    assert "no resume-builder" in text
    assert 'mode "score-only"' in text
    reviewer = (SKILLS_DIR / "ats" / "ats-scoring" / "SKILL.md").read_text()
    assert 'mode "score-only"' in reviewer and "`n/a`" in reviewer


def test_the_todo_list_is_the_only_plan():
    """No plan files anywhere: the todo list is the single source of truth."""
    planning = (ORCH / "workflow-planning" / "SKILL.md").read_text()
    assert "single source of truth" in planning and "write_todos" in planning
    assert "the user adds or changes something" in planning  # plan changes from user input
    for skill in SKILLS_DIR.glob("*/*/SKILL.md"):
        text = skill.read_text()
        assert "plan.md" not in text and "progress.md" not in text and "/output/plan" not in text, (
            skill
        )


def test_every_workflow_plans_with_todos():
    for name in ("jobhunt-workflow", "resume-optimization", "ats-score-only", "output-cleanup"):
        text = (ORCH / name / "SKILL.md").read_text().lower()
        assert "todo" in text, name


# -- cleanup ---------------------------------------------------------------------------

DELIVERABLES = [
    "resume/1-acme-backend/Jane Resume_acme-backend_redline.docx",
    "resume/1-acme-backend/Jane Resume_acme-backend_review.docx",
    "resume/1-acme-backend/Jane Resume_acme-backend_final.docx",
    "resume/2-beta-platform/Jane Resume_beta-platform_final.docx",
    "ats/ats_report.docx",
]
WORK_FILES = [
    "jobs/jobs.json",
    "jobs/search_notes.md",
    "match/match_report.json",
    "match/score_jobs.py",
    "scripts/dump_resume.py",
    ".DS_Store",
    "resume/1-acme-backend/revise_resume.py",
    "resume/1-acme-backend/change_log.md",
    "resume/1-acme-backend/ats_history.json",
    "resume/1-acme-backend/v1/Jane Resume_acme-backend_final.docx",  # round copy, not a deliverable
    "resume/1-acme-backend/v1/Jane Resume_acme-backend_review.docx",
    "resume/1-acme-backend/ats/round-1.json",
    "resume/1-acme-backend/ats/ats_score.py",
    "ats/general/round-1.json",
    "resume/1-acme-backend/~$ne Resume_acme-backend_final.docx",  # Word lock file
    ".browser/captures/page.png",  # browser output is agent output too
    ".browser/captures/listings_extract.json",
]
PROTECTED = [
    "original/Jane Resume.docx",
    ".browser/profile/Cookies",
    "memories/user_profile.md",
    "memories/agent_notes.md",
]


@pytest.fixture
def output(tmp_path):
    root = tmp_path / "output"
    for rel in DELIVERABLES + WORK_FILES + PROTECTED:
        path = root / rel
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(rel)
    (tmp_path / "cleanup.py").write_text(CLEANUP)
    return root


def cleanup(output, *flags):
    run = subprocess.run(
        ["python", str(output.parent / "cleanup.py"), "--root", str(output), *flags],
        capture_output=True,
        text=True,
        check=False,
    )
    assert run.returncode == 0, run.stderr
    return run.stdout


def files(root):
    return sorted(
        os.path.relpath(os.path.join(d, f), root) for d, _, fs in os.walk(root) for f in fs
    )


def test_dry_run_deletes_nothing_and_lists_the_plan(output):
    before = files(output)
    out = cleanup(output)
    assert files(output) == before
    keep = out.split("KEEP:")[1].split("DELETE")[0].split()
    assert "dry run" in out
    for rel in DELIVERABLES:
        assert rel.split("/")[-1].split()[-1] in " ".join(keep)
    assert f"DELETE ({len(WORK_FILES)} files)" in out


def test_apply_keeps_only_deliverables_and_protected(output):
    out = cleanup(output, "--apply")
    assert f"deleted {len(WORK_FILES)} files" in out
    assert files(output) == sorted(DELIVERABLES + PROTECTED)
    # emptied folders are gone, protected folders untouched
    for gone in (
        "jobs",
        "match",
        "scripts",
        "resume/1-acme-backend/v1",
        "resume/1-acme-backend/ats",
        "ats/general",
    ):
        assert not (output / gone).exists(), gone
    assert (output / ".browser" / "profile" / "Cookies").read_text() == ".browser/profile/Cookies"


def test_apply_is_idempotent(output):
    cleanup(output, "--apply")
    out = cleanup(output, "--apply")
    assert "deleted 0 files" in out
    assert files(output) == sorted(DELIVERABLES + PROTECTED)


def test_symlinks_are_not_followed_out_of_output(output, tmp_path):
    outside = tmp_path / "outside"
    outside.mkdir()
    (outside / "precious.docx").write_text("keep me")
    os.symlink(outside, output / "link-to-outside")
    cleanup(output, "--apply")
    assert (outside / "precious.docx").read_text() == "keep me"


def test_deliverable_patterns_are_strict():
    ns: dict = {}
    exec(CLEANUP.split('if __name__ == "__main__":')[0], ns)  # noqa: S102 - load the skill's code
    patterns = ns["DELIVERABLE"]
    ok = ["resume/1-x/cv_x_final.docx", "resume/2-y/cv_y_redline.docx", "ats/ats_report.docx"]
    bad = [
        "resume/1-x/v1/cv_x_final.docx",
        "resume/1-x/cv_x_final.docx.bak",
        "resume/cv_final.docx",
        "jobs/cv_final.docx",
        "ats/general/ats_report.docx",
        "resume/1-x/~$_x_final.docx",  # Word lock file
    ]
    assert all(any(p.match(r) for p in patterns) for r in ok)
    assert not any(any(p.match(r) for p in patterns) for r in bad)


# -- ATS-only report ---------------------------------------------------------------------


def test_ats_report_is_built_from_round_files(tmp_path):
    ats = tmp_path / "output" / "ats"
    rounds = {
        "acme-backend": {
            "score": 72.5,
            "job_title": "Backend Engineer",
            "breakdown": {
                "required_skills": 26.2,
                "keywords": 15.0,
                "sections": 7.5,
                "format": 10.5,
                "contact": 6.7,
                "nice_to_have": None,
            },
            "missing_required": ["Kafka"],
            "missing_keywords": ["microservices"],
            "issues": [
                {
                    "priority": "high",
                    "problem": "no standard 'summary' heading",
                    "quote": None,
                    "suggestion": "add a Summary",
                }
            ],
            "quality_notes": "Bullets lack numbers.",
        },
        "general": {
            "score": 64.0,
            "job_title": None,
            "breakdown": {"sections": 7.5, "required_skills": None},
            "issues": [],
        },
    }
    for slug, data in rounds.items():
        (ats / slug).mkdir(parents=True)
        (ats / slug / "round-1.json").write_text(json.dumps(data))
    script = REPORT.replace("/output/ats", str(ats))
    run = subprocess.run(["python", "-c", script], capture_output=True, text=True, check=False)
    assert run.returncode == 0, run.stderr
    report = docx.Document(io.BytesIO((ats / "ats_report.docx").read_bytes()))
    text = "\n".join(p.text for p in report.paragraphs)
    assert "Backend Engineer - score 72.5/100" in text
    assert "general - score 64.0/100" in text
    assert "Missing required skills: Kafka" in text
    assert "[high] no standard 'summary' heading -> add a Summary" in text
    cells = [c.text for t in report.tables for row in t.rows for c in row.cells]
    assert "n/a" in cells and "required skills" in cells


def test_skill_code_blocks_are_marked():
    for skill in ORCH.glob("*/SKILL.md"):
        for block in re.findall(
            r"^```python\n(.*?)\n```$", skill.read_text(), re.DOTALL | re.MULTILINE
        ):
            compile(block, str(skill), "exec")
