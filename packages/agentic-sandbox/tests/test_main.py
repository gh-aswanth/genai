"""Resume intake: validate, copy read-only, mount only the copy."""

from __future__ import annotations

import io
import os
import stat

import docx
import pytest
from genai_agentic_sandbox.builder import SandboxImageBuilder
from genai_agentic_sandbox.main import (
    MIN_RESUME_CHARS,
    SKILLS_DIR,
    ResumeError,
    docx_text_chars,
    main,
    prepare_resume,
    sandbox_mounts,
)

from .docx_samples import RESUME, render


def blank_docx() -> bytes:
    out = io.BytesIO()
    docx.Document().save(out)
    return out.getvalue()


@pytest.fixture
def resume(tmp_path):
    path = tmp_path / "home" / "Jane Resume.docx"
    path.parent.mkdir()
    path.write_bytes(RESUME)
    return path


def test_text_chars_counts_visible_text(resume):
    # Every w:t of the sample; tabs are w:tab elements, not text.
    expected = sum(len(line.replace("\t", "")) for line in render(RESUME))
    assert docx_text_chars(resume) == expected > MIN_RESUME_CHARS


def test_blank_document_is_rejected(tmp_path):
    """The failure from the real run: a blank .docx must stop the run up front."""
    blank = tmp_path / "resume.docx"
    blank.write_bytes(blank_docx())
    with pytest.raises(ResumeError, match="contains only 0 characters of text - it looks empty"):
        prepare_resume(blank, tmp_path / "out")
    assert not (tmp_path / "out" / "original").exists()


@pytest.mark.parametrize("name", ["resume.pdf", "missing.docx"])
def test_not_a_docx(tmp_path, name):
    path = tmp_path / name
    if name.endswith(".pdf"):
        path.write_bytes(b"%PDF")
    with pytest.raises(ResumeError, match="not a .docx"):
        prepare_resume(path, tmp_path / "out")


def test_corrupt_docx(tmp_path):
    path = tmp_path / "resume.docx"
    path.write_bytes(b"not a zip")
    with pytest.raises(ResumeError, match="not a readable .docx"):
        prepare_resume(path, tmp_path / "out")


def test_copy_is_identical_and_read_only(resume, tmp_path, monkeypatch):
    monkeypatch.setattr("genai_agentic_sandbox.main.MIN_RESUME_CHARS", 10)
    copy = prepare_resume(resume, tmp_path / "out")
    assert copy == tmp_path / "out" / "original" / "Jane Resume.docx"
    assert copy.read_bytes() == resume.read_bytes() == RESUME
    assert not copy.stat().st_mode & (stat.S_IWUSR | stat.S_IWGRP | stat.S_IWOTH)
    assert os.access(resume, os.W_OK)  # the user's own file is left as it was


def test_rerun_refreshes_the_copy(resume, tmp_path, monkeypatch):
    monkeypatch.setattr("genai_agentic_sandbox.main.MIN_RESUME_CHARS", 10)
    prepare_resume(resume, tmp_path / "out")
    edited = docx.Document(io.BytesIO(RESUME))
    edited.add_paragraph("New certification: AWS Solutions Architect")
    edited.save(resume)
    copy = prepare_resume(resume, tmp_path / "out")  # read-only copy is replaced
    assert copy.read_bytes() == resume.read_bytes()


def test_mounts_expose_only_the_copy(resume, tmp_path, monkeypatch):
    monkeypatch.setattr("genai_agentic_sandbox.main.MIN_RESUME_CHARS", 10)
    out = tmp_path / "out"
    copy = prepare_resume(resume, out)
    mounts = {m.container_path: m for m in sandbox_mounts(copy, out)}
    assert set(mounts) == {"/output", "/output/original", "/input", "/skills"}
    assert mounts["/output"].read_only is False
    for path in ("/output/original", "/input", "/skills"):
        assert mounts[path].read_only is True
    assert mounts["/input"].host_path == copy.parent
    assert mounts["/skills"].host_path == SKILLS_DIR
    assert all(
        resume.parent not in (m.host_path, *getattr(m.host_path, "parents", ()))
        for m in mounts.values()
    )


def test_cli_stops_on_empty_resume(tmp_path, monkeypatch, capsys):
    monkeypatch.setenv("OPENAI_API_KEY", "sk-test")
    blank = tmp_path / "resume.docx"
    blank.write_bytes(blank_docx())
    code = main(["--resume", str(blank), "--out", str(tmp_path / "out"), "find jobs"])
    assert code == 2
    assert "looks empty" in capsys.readouterr().err


@pytest.mark.integration
def test_container_cannot_modify_the_resume_copy(resume, tmp_path, monkeypatch):
    monkeypatch.setattr("genai_agentic_sandbox.main.MIN_RESUME_CHARS", 10)
    out = tmp_path / "out"
    copy = prepare_resume(resume, out)
    with SandboxImageBuilder().backend(mounts=sandbox_mounts(copy, out)) as sandbox:
        name = copy.name
        for target in (f"/input/{name}", f"/output/original/{name}"):
            result = sandbox.execute(f"python -c \"open('{target}', 'wb').write(b'x')\"")
            assert result.exit_code != 0, target
            assert sandbox.execute(f"rm -f '{target}'").exit_code != 0, target
        # ...while the rest of /output is writable and lands on the host.
        written = sandbox.execute(
            "mkdir -p /output/resume/1-acme && echo ok > /output/resume/1-acme/probe.txt"
        )
        assert written.exit_code == 0, written.output
        # The copy is readable where the agent is told to find it.
        size = sandbox.execute(f"python -c \"import os; print(os.path.getsize('/input/{name}'))\"")
    assert int(size.output.strip()) == len(RESUME)
    assert copy.read_bytes() == RESUME
    assert (out / "resume" / "1-acme" / "probe.txt").read_text() == "ok\n"


@pytest.mark.parametrize(
    "flags",
    [
        ["--top-jobs", "0"],
        ["--max-rounds", "0"],
        ["--target-score", "0"],
        ["--target-score", "101"],
    ],
)
def test_cli_rejects_bad_loop_settings(resume, tmp_path, monkeypatch, capsys, flags):
    monkeypatch.setenv("OPENAI_API_KEY", "sk-test")
    code = main(["--resume", str(resume), "--out", str(tmp_path / "out"), *flags, "find jobs"])
    assert code == 2
    assert "--max-rounds >= 1" in capsys.readouterr().err
