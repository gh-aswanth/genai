"""The ATS reviewer's reference scoring script (from its skill) must be sound.

It is what makes the builder <-> reviewer loop converge: deterministic, higher
for a better-targeted resume, and penalising what real ATS parsers choke on.
"""

from __future__ import annotations

import io
import json
import subprocess

import docx
import pytest
from genai_agentic_sandbox.main import SKILLS_DIR

from .docx_samples import RESUME, skill_code
from .test_skill_docx_code import script_for

SCORER = skill_code(
    (SKILLS_DIR / "ats" / "ats-scoring" / "SKILL.md").read_text(), "ats-score reference"
)
JOB = {
    "required_skills": ["Python", "FastAPI", "Kafka", "Docker"],
    "nice_to_have": ["AWS"],
    "keywords": ["REST", "APIs", "microservices", "latency"],
}
WEIGHTS = {
    "required_skills": 35,
    "keywords": 20,
    "sections": 15,
    "format": 15,
    "contact": 10,
    "nice_to_have": 5,
}


@pytest.fixture(scope="module")
def workdir(tmp_path_factory):
    d = tmp_path_factory.mktemp("ats")
    (d / "ats_score.py").write_text(SCORER)
    (d / "job.json").write_text(json.dumps(JOB))
    (d / "cv.docx").write_bytes(RESUME)
    build = d / "build.py"
    build.write_text(script_for(str(d / "cv.docx"), str(d / "job")))
    run = subprocess.run(["python", str(build)], capture_output=True, text=True, check=False)
    assert run.returncode == 0, run.stderr
    return d


def score(workdir, docx_path) -> dict:
    out = workdir / "out.json"
    run = subprocess.run(
        [
            "python",
            str(workdir / "ats_score.py"),
            str(docx_path),
            str(workdir / "job.json"),
            str(out),
        ],
        capture_output=True,
        text=True,
        check=False,
    )
    assert run.returncode == 0, run.stderr
    assert run.stdout.startswith("score ")
    return json.loads(out.read_text())


def save(doc, path):
    buf = io.BytesIO()
    doc.save(buf)
    path.write_bytes(buf.getvalue())
    return path


def test_breakdown_adds_up_within_weights(workdir):
    result = score(workdir, workdir / "cv.docx")
    assert set(result["breakdown"]) == set(WEIGHTS)
    for key, value in result["breakdown"].items():
        assert 0 <= value <= WEIGHTS[key]
    assert result["score"] == pytest.approx(sum(result["breakdown"].values()), abs=0.11)


def test_is_deterministic(workdir):
    assert score(workdir, workdir / "cv.docx") == score(workdir, workdir / "cv.docx")


def test_tailored_final_scores_higher_than_original(workdir):
    original = score(workdir, workdir / "cv.docx")
    final = score(workdir, workdir / "job" / "v1" / "cv_acme-backend_final.docx")
    assert final["score"] > original["score"]
    assert "FastAPI" in final["matched_required"] and "FastAPI" in original["missing_required"]
    assert final["missing_required"] == ["Kafka"]  # never invented, so still missing


def test_unaccepted_tracked_changes_are_penalised(workdir):
    final = score(workdir, workdir / "job" / "v1" / "cv_acme-backend_final.docx")
    review = score(workdir, workdir / "job" / "v1" / "cv_acme-backend_review.docx")
    assert review["breakdown"]["format"] < final["breakdown"]["format"]
    assert any("tracked changes" in i["problem"] for i in review["issues"])


def test_sections_and_issues(workdir):
    result = score(workdir, workdir / "cv.docx")
    assert result["sections"] == {
        "summary": False,
        "experience": True,
        "skills": True,
        "education": False,
    }
    problems = [i["problem"] for i in result["issues"]]
    assert "required skill 'Kafka' not found" in problems
    assert "no standard 'education' heading" in problems
    assert any("tables" in p for p in problems)  # the sample keeps contacts in a table
    assert all(i["priority"] in ("high", "medium", "low") for i in result["issues"])


def test_contact_in_header_is_penalised(workdir, tmp_path):
    d = docx.Document()
    d.sections[0].header.paragraphs[0].text = "jane@example.com | +91 98765 43210"
    for line in (
        "Summary",
        "Python engineer building REST APIs with FastAPI.",
        "Skills",
        "Python, FastAPI, Docker, Kafka",
        "Experience",
        "Acme - Engineer",
        "Education",
        "B.Tech",
    ):
        d.add_paragraph(line)
    result = score(workdir, save(d, tmp_path / "header.docx"))
    assert any("header/footer" in i["problem"] for i in result["issues"])
    assert result["contact"]["email"] is False  # not in the body, where ATS read it


def test_word_matching_is_whole_word(workdir, tmp_path):
    d = docx.Document()
    d.add_paragraph("Javascript developer; Reactive systems")
    (tmp_path / "job.json").write_text(json.dumps({"required_skills": ["Java", "React"]}))
    out = tmp_path / "o.json"
    subprocess.run(
        [
            "python",
            str(workdir / "ats_score.py"),
            str(save(d, tmp_path / "j.docx")),
            str(tmp_path / "job.json"),
            str(out),
        ],
        check=True,
        capture_output=True,
    )
    assert json.loads(out.read_text())["missing_required"] == ["Java", "React"]


def test_generic_check_without_a_job_marks_components_na(workdir, tmp_path):
    """Score-only with no job: keyword components are n/a, score is out of what applies."""
    (tmp_path / "job.json").write_text(
        json.dumps({"title": None, "required_skills": [], "keywords": []})
    )
    out = tmp_path / "o.json"
    subprocess.run(
        [
            "python",
            str(workdir / "ats_score.py"),
            str(workdir / "cv.docx"),
            str(tmp_path / "job.json"),
            str(out),
        ],
        check=True,
        capture_output=True,
    )
    result = json.loads(out.read_text())
    b = result["breakdown"]
    assert b["required_skills"] is None and b["keywords"] is None and b["nice_to_have"] is None
    applicable = WEIGHTS["sections"] + WEIGHTS["format"] + WEIGHTS["contact"]
    expected = 100 * (b["sections"] + b["format"] + b["contact"]) / applicable
    assert result["score"] == pytest.approx(expected, abs=0.11)
    assert not any(i["category"] in ("required_skill", "keyword") for i in result["issues"])


def test_job_title_is_reported(workdir, tmp_path):
    (tmp_path / "job.json").write_text(
        json.dumps({**JOB, "title": "Backend Engineer", "company": "Acme"})
    )
    out = tmp_path / "o.json"
    subprocess.run(
        [
            "python",
            str(workdir / "ats_score.py"),
            str(workdir / "cv.docx"),
            str(tmp_path / "job.json"),
            str(out),
        ],
        check=True,
        capture_output=True,
    )
    result = json.loads(out.read_text())
    assert (result["job_title"], result["company"]) == ("Backend Engineer", "Acme")
