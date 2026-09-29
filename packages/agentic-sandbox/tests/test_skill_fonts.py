"""New resume text must use the resume's own fonts (the bug from a real run).

The fixture mirrors that resume: body paragraphs use a style whose font differs
from the text (here "Body Web" = Times New Roman), and every run overrides it with
direct formatting (Georgia 10.5pt). Text added without copying a run's formatting
falls back to the style font - visibly different. The skill's helpers must always
copy the right run, and `check_fonts` must catch any text that does not.
"""

from __future__ import annotations

import io
import subprocess

import docx
import pytest
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt
from genai_agentic_sandbox.main import SKILLS_DIR

from .docx_samples import W, part, render, skill_code

HELPERS = skill_code(
    (SKILLS_DIR / "resume" / "docx-tracked-revisions" / "SKILL.md").read_text(),
    "tracked-change helpers",
)


def font_resume() -> bytes:
    d = docx.Document()
    body = d.styles.add_style("Body Web", WD_STYLE_TYPE.PARAGRAPH)
    body.font.name = "Times New Roman"
    body.font.size = Pt(12)

    def para(text, style="Body Web", bold=False):
        p = d.add_paragraph(style=style)
        run = p.add_run(text)
        run.font.name, run.font.size, run.bold = "Georgia", Pt(10.5), bold or None
        return p

    para("Jane Doe")  # 0
    para("Python developer building APIs.")  # 1
    para("Skills", bold=True)  # 2
    para("Python, Django, SQL")  # 3
    d.add_paragraph(style="Body Web")  # 4: empty paragraph (no runs)
    p = para("Portfolio: ")  # 5
    p.add_run("jane.dev").style = (
        d.styles["Hyperlink"] if "Hyperlink" in [s.name for s in d.styles] else None
    )
    out = io.BytesIO()
    d.save(out)
    return out.getvalue()


FONT_RESUME = font_resume()


def run_script(tmp_path, body: str) -> tuple[bytes, str]:
    src, out = tmp_path / "cv.docx", tmp_path / "out.docx"
    src.write_bytes(FONT_RESUME)
    script = tmp_path / "s.py"
    script.write_text(HELPERS + f"\nSRC, OUT = {str(src)!r}, {str(out)!r}\n" + body)
    run = subprocess.run(["python", str(script)], capture_output=True, text=True, check=False)
    assert run.returncode == 0, run.stderr
    return out.read_bytes(), run.stdout


def fonts_of(content: bytes, text: str) -> list[tuple[str | None, str | None]]:
    """(rFonts ascii, sz) of each run whose text is `text`."""
    out = []
    for r in part(content).iter(f"{W}r"):
        if "".join(t.text or "" for t in r.iter(f"{W}t")) == text:
            fonts, size = r.find(f"{W}rPr/{W}rFonts"), r.find(f"{W}rPr/{W}sz")
            out.append(
                (
                    fonts.get(f"{W}ascii") if fonts is not None else None,
                    size.get(f"{W}val") if size is not None else None,
                )
            )
    return out


GEORGIA = ("Georgia", "21")


def test_the_bug_is_caught(tmp_path):
    """Text added the old way (no template / python-docx add_run) gets the style font."""
    _, stdout = run_script(
        tmp_path,
        """
doc = open_doc(SRC)
P = paragraphs(doc)
ins = rev("w:ins"); ins.append(new_run(" and FastAPI", None)); P[1].append(ins)
from docx.text.paragraph import Paragraph
Paragraph(P[3], None).add_run(", Kafka")
doc.save(OUT)
print(check_fonts(SRC, OUT))
""",
    )
    assert "' and FastAPI', ('Times New Roman', '24')" in stdout
    assert "', Kafka', ('Times New Roman', '24')" in stdout


def test_every_helper_keeps_the_resume_font(tmp_path):
    content, stdout = run_script(
        tmp_path,
        """
doc = open_doc(SRC)
P = paragraphs(doc)
track_insert(P[1], " with FastAPI", after="building")
track_insert(P[1], " Remote OK.")
track_insert(P[3], "Go, ", after="")
track_replace(P[1], "Python developer", "Senior Python engineer")
insert_paragraph_after(P[3], "AWS, Docker")
insert_paragraph_after(P[0], "Summary: backend engineer.", like=P[1])
doc.save(OUT)
print("PROBLEMS", check_fonts(SRC, OUT))
""",
    )
    assert "PROBLEMS []" in stdout
    for text in (
        " with FastAPI",
        " Remote OK.",
        "Go, ",
        "Senior Python engineer",
        "AWS, Docker",
        "Summary: backend engineer.",
    ):
        assert fonts_of(content, text) == [GEORGIA], text


def test_empty_paragraph_borrows_from_a_same_style_paragraph(tmp_path):
    content, stdout = run_script(
        tmp_path,
        """
doc = open_doc(SRC)
P = paragraphs(doc)
track_insert(P[4], "Certifications: AWS SAA")
doc.save(OUT)
print("PROBLEMS", check_fonts(SRC, OUT))
""",
    )
    assert "PROBLEMS []" in stdout
    assert fonts_of(content, "Certifications: AWS SAA") == [GEORGIA]


def test_comment_reference_run_is_never_the_template(tmp_path):
    content, _ = run_script(
        tmp_path,
        """
doc = open_doc(SRC)
P = paragraphs(doc)
runs = visible_runs(P[3])
add_comment(doc, P[3], runs[0], runs[-1], "note")   # appends a commentReference run
track_insert(P[3], ", Kafka")                       # appended after it
doc.save(OUT)
""",
    )
    r = next(r for r in part(content).iter(f"{W}r") if "".join(r.itertext()) == ", Kafka")
    assert r.find(f"{W}rPr/{W}rStyle") is None  # not CommentReference (tiny superscript)
    assert fonts_of(content, ", Kafka") == [GEORGIA]


def test_review_colours_are_not_inherited(tmp_path):
    """In the review copy the neighbour may be flagged (pink) or recoloured: new text
    copies its ORIGINAL formatting, then gets its own review colour."""
    content, _ = run_script(
        tmp_path,
        """
doc = open_doc(SRC)
P = paragraphs(doc)
flag_wrong(doc, P[1], "building APIs", "building REST APIs", "vague")
track_insert(P[1], " (2020-2025)", after="building APIs")
doc.save(OUT)
""",
    )
    r = next(r for r in part(content).iter(f"{W}r") if "".join(r.itertext()) == " (2020-2025)")
    rpr = r.find(f"{W}rPr")
    assert rpr.find(f"{W}shd") is None and rpr.find(f"{W}u") is None  # no flag styling copied
    assert rpr.find(f"{W}color").get(f"{W}val") == "00B050"  # its own colour: added
    assert fonts_of(content, " (2020-2025)") == [GEORGIA]


def test_theme_fonts_are_resolved(tmp_path):
    (tmp_path / "t.docx").write_bytes(FONT_RESUME)
    script = tmp_path / "s.py"
    script.write_text(
        HELPERS + f"\nprint(_theme_fonts({str(tmp_path / 't.docx')!r}).get('minorHAnsi'))\n"
    )
    out = subprocess.run(["python", str(script)], capture_output=True, text=True, check=True).stdout
    assert out.strip() not in ("", "None")  # e.g. "Calibri"


def test_check_is_in_the_example_and_rules():
    skill = (SKILLS_DIR / "resume" / "docx-tracked-revisions" / "SKILL.md").read_text()
    assert 'print("CHECK fonts match the original:", not font_problems)' in skill
    assert "add_run" in skill and "template_run" in skill


@pytest.mark.parametrize("view", ["original", "accepted"])
def test_font_fixture_is_sane(view):
    assert render(FONT_RESUME, view)[1] == "Python developer building APIs."


def test_styles_the_resume_never_uses_are_refused(tmp_path):
    """The second bug from the real run: new bullets styled "ListBullet" when the resume's
    bullets use another style - they took the template's default font."""
    _, stdout = run_script(
        tmp_path,
        """
doc = open_doc(SRC)
P = paragraphs(doc)
for fn in (lambda: insert_paragraph_after(P[3], "New bullet", style_id="ListBullet"),
           lambda: track_paragraph_style(P[2], "Heading2")):
    try:
        fn()
        print("ALLOWED")
    except ValueError as exc:
        print("REFUSED", exc)
insert_paragraph_after(P[3], "Like an existing one", like=P[3])   # the right way
doc.save(OUT)
print("PROBLEMS", check_fonts(SRC, OUT))
""",
    )
    assert stdout.count("REFUSED") == 2 and "ALLOWED" not in stdout
    assert "is not used in this resume" in stdout and "BodyWeb" in stdout
    assert "PROBLEMS []" in stdout
